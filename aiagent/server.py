import json
import os
import uuid
from copy import deepcopy
from datetime import datetime
from io import BytesIO
from pathlib import Path
from urllib import error, request

from flask import Flask, jsonify, request as flask_request
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent
CONFIG_PATH = ROOT / "config.py"

if not CONFIG_PATH.exists():
    CONFIG_PATH.write_text((ROOT / "config.example.py").read_text(encoding="utf-8"), encoding="utf-8")

config = {}
exec(CONFIG_PATH.read_text(encoding="utf-8"), config)

API_KEY = (config.get("API_KEY") or "").strip()
MODEL = (config.get("MODEL") or "").strip()
BASE_URL = (config.get("BASE_URL") or "https://api.groq.com/openai/v1").rstrip("/")
TEMPERATURE = config.get("TEMPERATURE", 1)
TOP_P = config.get("TOP_P", 1)
MAX_COMPLETION_TOKENS = config.get("MAX_COMPLETION_TOKENS", 2048)
REASONING_EFFORT = (config.get("REASONING_EFFORT") or "").strip()
LMS_BASE = (config.get("LMS_BASE") or "http://127.0.0.1/LMS").rstrip("/")
REPORT_URL = LMS_BASE + "/aiagent/data/report.php"

CONTEXT = (ROOT / "source_context.md").read_text(encoding="utf-8")
SYSTEM = (
    CONTEXT
    + "\n\nThe owner is logged into Link Mark. Speak in plain sentences. "
    + "The table on screen is the data. Your reply is only the explanation."
)

app = Flask(__name__)
SESSIONS = {}
LAST_LIMITS = {"model": MODEL}

TOOLS = [
    {
        "type": "function",
        "function": {
            "name": "get_report",
            "description": "Load a full Link Mark report. Leave filters empty unless the owner asked for them.",
            "parameters": {
                "type": "object",
                "properties": {
                    "report": {
                        "type": "string",
                        "enum": [
                            "purchase",
                            "payable",
                            "stock",
                            "general_ledger",
                            "packing_warehouse",
                            "packing_gatepass",
                            "profit_and_loss",
                        ],
                    },
                    "date_from": {"type": "string"},
                    "date_to": {"type": "string"},
                    "supplier": {"type": "string"},
                    "commodity": {"type": "string"},
                    "voucher_no": {"type": "string"},
                    "size": {"type": "string"},
                    "supplier_type": {"type": "string"},
                    "stock_view": {"type": "string"},
                    "direction": {"type": "string"},
                    "country": {"type": "string"},
                    "fish_type": {"type": "string"},
                    "account_code": {"type": "string"},
                    "material": {"type": "string"},
                    "destination": {"type": "string"},
                    "movement": {"type": "string"},
                },
                "required": ["report"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "shape_table",
            "description": "Filter, sort, choose columns, or group the report already on screen. Every matching row is kept.",
            "parameters": {
                "type": "object",
                "properties": {
                    "table_id": {"type": "string"},
                    "columns": {"type": "array", "items": {"type": "string"}},
                    "filters": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "column": {"type": "string"},
                                "op": {"type": "string", "enum": ["eq", "contains", "gt", "gte", "lt", "lte"]},
                                "value": {"type": "string"},
                            },
                            "required": ["column", "op", "value"],
                        },
                    },
                    "sort_by": {"type": "string"},
                    "sort_dir": {"type": "string", "enum": ["asc", "desc"]},
                    "group_by": {"type": "string"},
                    "sum_columns": {"type": "array", "items": {"type": "string"}},
                },
            },
        },
    },
]


def owner_name():
    raw = flask_request.headers.get("X-LMS-User", "owner")
    safe = "".join(ch for ch in raw if ch.isalnum() or ch in "-_")
    return safe or "owner"


def chat_folder():
    folder = os.path.join(os.path.dirname(__file__), "chats", owner_name())
    os.makedirs(folder, exist_ok=True)
    return folder


def chat_file(number):
    return os.path.join(chat_folder(), f"chat{int(number)}.txt")


def now_stamp():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def write_record(record):
    record["updated"] = now_stamp()
    if record.get("title") in (None, "", "New chat"):
        for turn in record.get("turns", []):
            if turn.get("role") == "user" and turn.get("text"):
                record["title"] = " ".join(turn["text"].split())[:48]
                break
        record["title"] = record.get("title") or "New chat"
    with open(chat_file(record["number"]), "w", encoding="utf-8") as handle:
        json.dump(record, handle, ensure_ascii=False, indent=2)


def read_record(number):
    path = chat_file(number)
    if not os.path.isfile(path):
        return None
    with open(path, encoding="utf-8") as handle:
        return json.load(handle)


def next_chat_number():
    highest = 0
    for name in os.listdir(chat_folder()):
        if name.startswith("chat") and name.endswith(".txt") and name[4:-4].isdigit():
            highest = max(highest, int(name[4:-4]))
    return highest + 1


def create_record():
    record = {"number": next_chat_number(), "title": "New chat", "updated": now_stamp(), "turns": []}
    write_record(record)
    return record


def state_for(record):
    messages = [{"role": "system", "content": SYSTEM}]
    tables = {}
    last_table_id = None
    for turn in record.get("turns", []):
        text = turn.get("text") or ""
        if turn.get("role") == "user":
            messages.append({"role": "user", "content": text})
        else:
            messages.append({"role": "assistant", "content": text or "Here is the table."})
        for table in turn.get("tables") or []:
            tables[table["id"]] = table
            last_table_id = table["id"]
    return {"messages": messages, "tables": tables, "last_table_id": last_table_id, "record": record}


def open_chat(number):
    record = read_record(number) if number else None
    if record is None:
        record = create_record()
    state = state_for(record)
    SESSIONS[f"{owner_name()}:{record['number']}"] = state
    return state


def remember_turn(state, role, text, tables=None, debug=None):
    state["record"]["turns"].append({
        "role": role,
        "text": text or "",
        "tables": tables or [],
        "debug": debug or [],
    })
    if role == "user":
        state["messages"].append({"role": "user", "content": text or ""})
    write_record(state["record"])


def chat_summaries():
    found = []
    for name in os.listdir(chat_folder()):
        if not (name.startswith("chat") and name.endswith(".txt")):
            continue
        with open(os.path.join(chat_folder(), name), encoding="utf-8") as handle:
            record = json.load(handle)
        found.append({
            "number": record.get("number"),
            "title": record.get("title") or "New chat",
            "updated": record.get("updated") or "",
        })
    found.sort(key=lambda item: item["number"] or 0, reverse=True)
    return found


def numeric_totals(columns, rows):
    totals = {}
    for column in columns:
        key = column["key"]
        if key in ("balance",):
            continue
        values = []
        for row in rows:
            value = row.get(key)
            if isinstance(value, bool):
                continue
            if isinstance(value, (int, float)):
                values.append(float(value))
        if values and len(values) == len(rows):
            totals[key] = round(sum(values), 3)
    return totals


def store_table(state, table):
    table = deepcopy(table)
    table["id"] = uuid.uuid4().hex
    table["totals"] = numeric_totals(table.get("columns", []), table.get("rows", []))
    state["tables"][table["id"]] = table
    state["last_table_id"] = table["id"]
    return table


def model_view(table):
    return {
        "table_id": table["id"],
        "title": table.get("title"),
        "row_count": table.get("row_count", len(table.get("rows", []))),
        "columns": table.get("columns", []),
        "totals": table.get("totals", {}),
        "filters": table.get("filters", {}),
        "note": "Every row is already in the owner's table. Do not list rows or say a sample is showing.",
    }


def call_report(cookie, arguments):
    payload = json.dumps(arguments).encode("utf-8")
    req = request.Request(
        REPORT_URL,
        data=payload,
        headers={"Content-Type": "application/json", "Cookie": cookie},
        method="POST",
    )
    try:
        with request.urlopen(req, timeout=120) as response:
            return json.loads(response.read().decode("utf-8"))
    except error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        try:
            parsed = json.loads(detail)
            message = parsed.get("error", detail)
        except json.JSONDecodeError:
            message = detail or exc.reason
        return {"error": message}


def compare(left, op, right):
    if op == "contains":
        return str(right).lower() in str(left).lower()
    if op == "eq":
        return str(left).lower() == str(right).lower()
    try:
        left_num = float(left)
        right_num = float(right)
    except (TypeError, ValueError):
        left_num = str(left)
        right_num = str(right)
    if op == "gt":
        return left_num > right_num
    if op == "gte":
        return left_num >= right_num
    if op == "lt":
        return left_num < right_num
    if op == "lte":
        return left_num <= right_num
    return False


def shape_table(state, arguments):
    table_id = arguments.get("table_id") or state.get("last_table_id")
    source = state["tables"].get(table_id)
    if not source:
        return {"error": "Load a report before shaping it."}
    rows = list(source["rows"])
    for item in arguments.get("filters") or []:
        column = item.get("column")
        op = item.get("op")
        value = item.get("value")
        rows = [row for row in rows if compare(row.get(column), op, value)]

    columns = source["columns"]
    keep = arguments.get("columns") or []
    if keep:
        columns = [column for column in columns if column["key"] in keep]
        if not columns:
            columns = [{"key": key, "label": key} for key in keep]

    group_by = arguments.get("group_by") or ""
    if group_by:
        sum_columns = arguments.get("sum_columns") or [
            column["key"] for column in columns if column["key"] != group_by
        ]
        grouped = {}
        for row in rows:
            key = str(row.get(group_by, ""))
            if key not in grouped:
                grouped[key] = {"count": 0}
            grouped[key]["count"] += 1
            for name in sum_columns:
                try:
                    grouped[key][name] = grouped[key].get(name, 0) + float(row.get(name) or 0)
                except (TypeError, ValueError):
                    grouped[key][name] = row.get(name)
        rows = []
        for key, values in grouped.items():
            item = {group_by: key, "count": values["count"]}
            for name in sum_columns:
                if name == group_by:
                    continue
                item[name] = values.get(name, 0)
            rows.append(item)
        columns = [{"key": group_by, "label": group_by}]
        columns.append({"key": "count", "label": "Count"})
        for name in sum_columns:
            if name != group_by:
                columns.append({"key": name, "label": name})

    sort_by = arguments.get("sort_by") or ""
    if sort_by:
        reverse = (arguments.get("sort_dir") or "asc") == "desc"

        def sort_key(row):
            value = row.get(sort_by)
            try:
                return (0, float(value))
            except (TypeError, ValueError):
                return (1, str(value).lower())

        rows.sort(key=sort_key, reverse=reverse)

    shaped = {
        "title": source["title"] + " (shaped)",
        "columns": columns,
        "rows": rows,
        "row_count": len(rows),
        "filters": {"shaped_from": table_id, "request": arguments},
    }
    return {"tables": [shaped]}


def run_tool(state, cookie, name, arguments):
    if name == "get_report":
        result = call_report(cookie, arguments)
    elif name == "shape_table":
        result = shape_table(state, arguments)
    else:
        result = {"error": "Unknown tool"}
    if result.get("error"):
        return result, []
    stored = []
    views = []
    for table in result.get("tables", []):
        saved = store_table(state, table)
        stored.append(saved)
        views.append(model_view(saved))
    return {"tables": views}, stored


def header_value(headers, name):
    if headers is None:
        return None
    return headers.get(name)


def remember_limits(headers, usage=None):
    info = {
        "model": MODEL,
        "limit_tokens": header_value(headers, "x-ratelimit-limit-tokens"),
        "remaining_tokens": header_value(headers, "x-ratelimit-remaining-tokens"),
        "reset_tokens": header_value(headers, "x-ratelimit-reset-tokens"),
        "limit_requests": header_value(headers, "x-ratelimit-limit-requests"),
        "remaining_requests": header_value(headers, "x-ratelimit-remaining-requests"),
        "reset_requests": header_value(headers, "x-ratelimit-reset-requests"),
    }
    if usage:
        info["prompt_tokens"] = usage.get("prompt_tokens")
        info["completion_tokens"] = usage.get("completion_tokens")
        info["total_tokens"] = usage.get("total_tokens")
    for key, value in info.items():
        if value not in (None, ""):
            LAST_LIMITS[key] = value
    return dict(LAST_LIMITS)


def asks_for_rows(text):
    lowered = text.lower()
    question = any(part in lowered for part in ("what ", "which ", "list", "can i", "how do"))
    show = any(part in lowered for part in ("show", "give me", "load", "open the", "export", "download"))
    if question and not show:
        return False
    return any(part in lowered for part in (
        "show", "give me", "load", "purchase", "payable", "statement", "stock",
        "ledger", "warehouse", "gate pass", "profit", "table", "supplier",
    ))


def llm(messages, tool_choice=None):
    if not API_KEY or not MODEL:
        raise RuntimeError("Open aiagent/config.py and set API_KEY and MODEL, then restart start.bat.")
    body = {
        "model": MODEL,
        "temperature": TEMPERATURE,
        "top_p": TOP_P,
        "max_completion_tokens": MAX_COMPLETION_TOKENS,
        "messages": messages,
        "tools": TOOLS,
        "stream": False,
    }
    if tool_choice:
        body["tool_choice"] = tool_choice
    if REASONING_EFFORT and "gpt-oss" in MODEL:
        body["reasoning_effort"] = REASONING_EFFORT
    payload = json.dumps(body).encode("utf-8")
    req = request.Request(
        BASE_URL + "/chat/completions",
        data=payload,
        headers={
            "Authorization": "Bearer " + API_KEY,
            "Content-Type": "application/json",
            "User-Agent": "LinkMarkAgent/1.0",
        },
        method="POST",
    )
    try:
        with request.urlopen(req, timeout=120) as response:
            data = json.loads(response.read().decode("utf-8"))
            remember_limits(response.headers, data.get("usage"))
            return data
    except error.HTTPError as exc:
        remember_limits(exc.headers)
        detail = exc.read().decode("utf-8", errors="replace")
        message = detail or str(exc.reason)
        try:
            parsed = json.loads(detail)
            err = parsed.get("error") or {}
            if err.get("code") == "rate_limit_exceeded":
                wait = ""
                text = err.get("message") or ""
                marker = "Please try again in "
                if marker in text:
                    wait = text.split(marker, 1)[1].split(".", 1)[0].strip()
                message = "Groq rate limit on this model. Wait " + (wait or "a few seconds") + " and send it again."
            elif err.get("message"):
                message = err["message"]
        except json.JSONDecodeError:
            pass
        raise RuntimeError(message) from exc


@app.post("/chat")
def chat():
    body = flask_request.get_json(force=True, silent=True) or {}
    message = (body.get("message") or "").strip()
    if not message:
        return jsonify({"error": "Type a message first."}), 400
    state = open_chat(int(body.get("chat_id") or 0))
    cookie = flask_request.headers.get("Cookie", "")
    remember_turn(state, "user", message)
    chat_number = state["record"]["number"]
    shown = []
    debug = []

    def note(step, detail):
        debug.append({"step": step, "detail": str(detail)})

    note("request", message)
    note("model", MODEL)

    def consume(tool_choice=None):
        note("call", "tool_choice=" + (tool_choice or "auto"))
        data = llm(state["messages"], tool_choice)
        choice = data["choices"][0]["message"]
        usage = data.get("usage") or {}
        note(
            "tokens",
            "prompt {prompt} + completion {completion} = {total}. Remaining {remaining} / {limit}. Reset {reset}.".format(
                prompt=usage.get("prompt_tokens", "?"),
                completion=usage.get("completion_tokens", "?"),
                total=usage.get("total_tokens", "?"),
                remaining=LAST_LIMITS.get("remaining_tokens", "?"),
                limit=LAST_LIMITS.get("limit_tokens", "?"),
                reset=LAST_LIMITS.get("reset_tokens", "?"),
            ),
        )
        tool_calls = choice.get("tool_calls") or []
        finish = data.get("choices", [{}])[0].get("finish_reason")
        note("finish", finish or "unknown")
        if choice.get("content"):
            note("reply", choice.get("content")[:400])
        assistant = {"role": "assistant", "content": choice.get("content") or None}
        if choice.get("reasoning"):
            assistant["reasoning"] = choice["reasoning"]
        if tool_calls:
            assistant["tool_calls"] = tool_calls
        state["messages"].append(assistant)
        if not tool_calls:
            return False
        for call in tool_calls:
            name = call.get("function", {}).get("name")
            raw_args = call.get("function", {}).get("arguments") or "{}"
            if isinstance(raw_args, dict):
                arguments = raw_args
            else:
                try:
                    arguments = json.loads(raw_args)
                except json.JSONDecodeError:
                    arguments = {}
            view, tables = run_tool(state, cookie, name, arguments)
            shown.extend(tables)
            if view.get("error"):
                note("tool", name + " failed: " + str(view.get("error")))
            else:
                counts = ", ".join(
                    str(item.get("title")) + " " + str(item.get("row_count")) + " rows" for item in view.get("tables", [])
                )
                note("tool", name + " " + json.dumps(arguments) + " -> " + (counts or "no rows"))
            state["messages"].append(
                {
                    "role": "tool",
                    "tool_call_id": call.get("id", "call"),
                    "content": json.dumps(view),
                }
            )
        return True

    try:
        for _ in range(6):
            if not consume():
                break
        if not shown and asks_for_rows(message):
            note("retry", "No table yet, so the next call requires a tool.")
            state["messages"].append(
                {
                    "role": "user",
                    "content": "No table was loaded. Call get_report for the report the owner asked for. Leave filters empty unless they named one. Do not invent rows or a row count.",
                }
            )
            try:
                for _ in range(2):
                    called = consume("required")
                    if not called or shown:
                        break
            except Exception as exc:
                note("error", str(exc))
                if "did not call a tool" not in str(exc).lower():
                    raise
        elif not shown:
            note("skip", "Question did not ask for rows, so no tool was required.")
    except Exception as exc:
        note("error", str(exc))
        remember_turn(state, "assistant", str(exc), [], debug)
        return jsonify({
            "error": str(exc),
            "debug": debug,
            "limits": LAST_LIMITS,
            "chat_id": chat_number,
            "title": state["record"]["title"],
        }), 500

    reply = ""
    for past in reversed(state["messages"]):
        if past.get("role") == "assistant" and past.get("content"):
            reply = past["content"]
            break
    if not reply and shown:
        reply = "Here is the full table."
    if not reply and not shown:
        reply = "I can open Purchase, Payable, Stock, General Ledger, Packing Material Warehouse, Gate Pass, and Profit and Loss."

    public_tables = []
    for table in shown:
        public = deepcopy(table)
        public_tables.append(public)
    remember_turn(state, "assistant", reply, public_tables, debug)
    return jsonify({
        "reply": reply,
        "tables": public_tables,
        "debug": debug,
        "limits": LAST_LIMITS,
        "chat_id": chat_number,
        "title": state["record"]["title"],
    })


def find_table(table_id, chat_id):
    record = read_record(int(chat_id or 0))
    if not record:
        return None
    for turn in record.get("turns", []):
        for table in turn.get("tables") or []:
            if table.get("id") == table_id:
                return table
    return None


@app.post("/export")
def export():
    body = flask_request.get_json(force=True, silent=True) or {}
    table = find_table(body.get("table_id", ""), body.get("chat_id"))
    if not table:
        return jsonify({"error": "That table is no longer available. Ask for the report again."}), 404
    kind = body.get("format")
    if kind == "xlsx":
        content = build_xlsx(table)
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        filename = "report.xlsx"
    elif kind == "pdf":
        content = build_pdf(table)
        mime = "application/pdf"
        filename = "report.pdf"
    else:
        return jsonify({"error": "Use xlsx or pdf."}), 400
    return app.response_class(content, mimetype=mime, headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def build_xlsx(table):
    book = Workbook()
    sheet = book.active
    sheet.title = "Report"
    headers = [column["label"] for column in table["columns"]]
    sheet.append([table["title"]])
    sheet.append(headers)
    for cell in sheet[2]:
        cell.font = Font(bold=True)
    keys = [column["key"] for column in table["columns"]]
    for row in table["rows"]:
        sheet.append([row.get(key, "") for key in keys])
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def build_pdf(table):
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=10)
    pdf.add_page()
    def pdf_text(value):
        text = str(value)
        return text.encode("latin-1", "replace").decode("latin-1")

    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, pdf_text(table["title"])[:120], new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 7)
    pdf.cell(0, 6, f"All {len(table['rows'])} rows", new_x="LMARGIN", new_y="NEXT")
    columns = table["columns"]
    width = max(18, min(45, 270 / max(len(columns), 1)))
    pdf.set_font("Helvetica", "B", 7)
    for column in columns:
        pdf.cell(width, 6, pdf_text(column["label"])[:22], border=1)
    pdf.ln()
    pdf.set_font("Helvetica", "", 7)
    for row in table["rows"]:
        for column in columns:
            value = row.get(column["key"], "")
            if isinstance(value, float):
                value = round(value, 3)
            pdf.cell(width, 5, pdf_text(value)[:22], border=1)
        pdf.ln()
    return bytes(pdf.output())


@app.get("/chats")
def list_chats():
    return jsonify({"chats": chat_summaries()})


@app.post("/chats")
def new_chat():
    record = create_record()
    return jsonify({
        "number": record["number"],
        "title": record["title"],
        "updated": record["updated"],
        "turns": [],
    })


@app.get("/chats/<int:number>")
def one_chat(number):
    record = read_record(number)
    if record is None:
        return jsonify({"error": "That chat does not exist."}), 404
    return jsonify(record)


@app.get("/health")
def health():
    return jsonify({"ok": True, "model_set": bool(MODEL), "key_set": bool(API_KEY)})


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=8765, debug=False)
